import io
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import sqlite3
import streamlit as st


# --- FUNGSI BANTU PENCOCOKAN & PARSING ---
def clean_tes_name(name):
  if not name:
    return ""
  cleaned = re.sub(
      r"\b(sts|pas|pts|uh|us|tp|uas|uts)\b", "", str(name), flags=re.IGNORECASE
  )
  return re.sub(r"\s+", " ", cleaned).strip().lower()


def clean_name(name):
  if not name:
    return ""
  return re.sub(r"\s+", " ", str(name)).strip().lower()


def parse_poin(val):
  if pd.isna(val):
    return 0.0
  if isinstance(val, (int, float)):
    return float(val)
  val_str = str(val).strip().replace(",", ".")
  try:
    return float(val_str)
  except:
    return 0.0


# --- KONEKSI DATABASE SQLITE & MIGRASI OTOMATIS ---
def init_db():
  conn = sqlite3.connect("sekolah.db", check_same_thread=False)
  cursor = conn.cursor()

  # Tabel Master Siswa
  cursor.execute("""
        CREATE TABLE IF NOT EXISTS master_siswa (
            nis TEXT PRIMARY KEY,
            nama TEXT,
            kelas TEXT
        )
    """)

  # Tabel Hasil Ujian Mentah
  cursor.execute("""
        CREATE TABLE IF NOT EXISTS hasil_ujian (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nama_tes TEXT,
            nis TEXT,
            nama TEXT,
            kelas TEXT,
            poin REAL,
            status TEXT DEFAULT 'normal',
            UNIQUE(nis, nama_tes)
        )
    """)

  # Tabel Data Remedial Siswa
  cursor.execute("""
        CREATE TABLE IF NOT EXISTS remedial_siswa (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nama_tes TEXT,
            nis TEXT,
            nama TEXT,
            kelas TEXT
        )
    """)
  conn.commit()
  return conn


conn = init_db()

st.set_page_config(page_title="APLIKASI REKAP NILAI UJIAN", layout="wide")

# --- KELOLA SESI LOGIN ---
if "logged_in" not in st.session_state:
  st.session_state.logged_in = False
  st.session_state.role = None

if not st.session_state.logged_in:
  st.title("🔐 Silakan Login ke Sistem")
  st.markdown("Masukkan akun akses sesuai peran Anda.")

  with st.form("login_form"):
    username = st.text_input("Username")
    password = st.text_input("Password", type="password")
    submit = st.form_submit_button("Login")

    if submit:
      if username == "admin" and password == "123":
        st.session_state.logged_in = True
        st.session_state.role = "admin"
        st.success("Login berhasil sebagai Admin!")
        st.rerun()
      elif username == "guru" and password == "123":
        st.session_state.logged_in = True
        st.session_state.role = "guru"
        st.success("Login berhasil sebagai Guru Mapel!")
        st.rerun()
      else:
        st.error("Username atau Password salah!")

else:
  st.sidebar.write(
      f"👤 Login sebagai: **{st.session_state.role.upper()}**"
  )
  if st.sidebar.button("🚪 Logout"):
    st.session_state.logged_in = False
    st.session_state.role = None
    st.rerun()

  st.title("APLIKASI REKAP NILAI UJIAN")
  st.markdown("---")

  if st.session_state.role == "admin":
    menu = st.sidebar.selectbox(
        "Pilih Menu:",
        [
            "1. Operator: Upload & Kelola Data",
            "2. Guru: Download Rekap Nilai",
            "3. Cek Siswa Belum Ujian",
            "4. Cek & Download Peserta Remedial",
        ],
    )
  else:
    menu = st.sidebar.selectbox(
        "Pilih Menu:",
        [
            "2. Guru: Download Rekap Nilai",
            "3. Cek Siswa Belum Ujian",
            "4. Cek & Download Peserta Remedial",
        ],
    )

  # --- MENU 1: OPERATOR UPLOAD & KELOLA ---
  if menu == "1. Operator: Upload & Kelola Data":
    st.header("Panel Operator: Upload & Kelola Database")

    tab_up1, tab_up2, tab_up3, tab_up4, tab_up5 = st.tabs([
        "Upload Master Siswa",
        "Upload Hasil Ujian CBT",
        "Upload Data Remedial",
        "Update Nilai & Selesaikan Remedial",
        "Kelola / Hapus Data",
    ])

    with tab_up1:
      st.subheader("1. Upload Data Master Siswa (Excel/CSV)")
      file_master = st.file_uploader(
          "Pilih file Master Siswa", type=["xlsx", "csv"], key="master"
      )
      if file_master:
        df_m = (
            pd.read_csv(file_master)
            if file_master.name.endswith(".csv")
            else pd.read_excel(file_master)
        )
        df_m.columns = [str(col).strip() for col in df_m.columns]
        col_map = {c.lower(): c for c in df_m.columns}
        nis_col = col_map.get("nis") or col_map.get("no induk")
        nama_col = col_map.get("nama") or col_map.get("nama lengkap")
        kelas_col = col_map.get("kelas") or col_map.get("group")

        if nis_col and nama_col and kelas_col:
          st.markdown("**Pratinjau Data Master:**")
          st.dataframe(
              df_m.head(), hide_index=True, use_container_width=True
          )
          if st.button("Simpan Master Siswa ke Database"):
            cursor = conn.cursor()
            for _, row in df_m.iterrows():
              nis_raw = row[nis_col]
              nis_clean = (
                  str(int(nis_raw))
                  if isinstance(nis_raw, float)
                  else str(nis_raw).strip()
              )
              cursor.execute(
                  "INSERT OR REPLACE INTO master_siswa (nis, nama, kelas)"
                  " VALUES (?, ?, ?)",
                  (
                      nis_clean,
                      str(row[nama_col]).strip(),
                      str(row[kelas_col]).strip(),
                  ),
              )
            conn.commit()
            st.success("Data Master Siswa berhasil disimpan!")

    with tab_up2:
      st.subheader("2. Upload Hasil Ujian CBT (Utama / Susulan)")
      jenis_upload = st.radio(
          "Pilih Jenis Upload File Ujian:",
          [
              "Ujian Utama CBT (Normal)",
              "Ujian Susulan / Susulan & Remedial",
          ],
      )

      file_ujian = st.file_uploader(
          "Pilih file Hasil Ujian", type=["xlsx", "csv"], key="ujian"
      )
      if file_ujian:
        df_u = (
            pd.read_csv(file_ujian)
            if file_ujian.name.endswith(".csv")
            else pd.read_excel(file_ujian)
        )
        df_u.columns = [str(col).strip() for col in df_u.columns]
        col_map_u = {c.lower(): c for c in df_u.columns}
        tes_c = col_map_u.get("nama tes") or col_map_u.get("namates")
        nis_c = (
            col_map_u.get("username")
            or col_map_u.get("nis")
            or col_map_u.get("no induk")
        )
        nama_c = col_map_u.get("nama")
        kelas_c = col_map_u.get("group") or col_map_u.get("kelas")
        poin_c = col_map_u.get("poin") or col_map_u.get("nilai")

        if tes_c and nama_c and poin_c:
          st.markdown("**Pratinjau Data Ujian:**")
          st.dataframe(
              df_u.head(), hide_index=True, use_container_width=True
          )
          if st.button("Simpan Hasil Ujian ke Database"):
            cursor = conn.cursor()
            for _, row in df_u.iterrows():
              nama_val = str(row[nama_c]).strip()
              nama_tes_val = str(row[tes_c]).strip()
              clean_upload_tes = clean_tes_name(nama_tes_val)
              poin_val = parse_poin(row[poin_c])

              nis_val = None
              if nis_c and pd.notna(row[nis_c]):
                nis_raw = row[nis_c]
                nis_val = (
                    str(int(nis_raw))
                    if isinstance(nis_raw, float)
                    else str(nis_raw).strip()
                )

              kelas_val = (
                  str(row[kelas_c]).strip()
                  if kelas_c and pd.notna(row[kelas_c])
                  else "-"
              )

              cursor.execute(
                  "SELECT nis, kelas FROM master_siswa WHERE LOWER(nama) ="
                  " LOWER(?)",
                  (nama_val,),
              )
              res_m = cursor.fetchone()
              if res_m:
                if not nis_val or nis_val.lower() == "nan":
                  nis_val = res_m[0]
                if kelas_val == "-" or not kelas_val:
                  kelas_val = res_m[1]

              if not nis_val or nis_val.lower() == "nan":
                nis_val = "UNKNOWN_" + nama_val

              if "Utama" in jenis_upload:
                status_ujian = "normal"
              else:
                cursor.execute(
                    "SELECT nama_tes FROM remedial_siswa WHERE LOWER(nama) ="
                    " LOWER(?)",
                    (nama_val,),
                )
                all_rem = cursor.fetchall()
                is_remedial = False
                for (r_tes,) in all_rem:
                  clean_r_tes = clean_tes_name(r_tes)
                  if (
                      clean_r_tes == clean_upload_tes
                      or clean_r_tes in clean_upload_tes
                      or clean_upload_tes in clean_r_tes
                  ):
                    is_remedial = True
                    break

                if is_remedial:
                  status_ujian = "remedial"
                else:
                  status_ujian = "susulan"

              cursor.execute(
                  """
                                INSERT OR REPLACE INTO hasil_ujian (nama_tes, nis, nama, kelas, poin, status)
                                VALUES (?, ?, ?, ?, ?, ?)
                            """,
                  (
                      nama_tes_val,
                      nis_val,
                      nama_val,
                      kelas_val,
                      poin_val,
                      status_ujian,
                  ),
              )
            conn.commit()
            st.success("Data Hasil Ujian berhasil di-upload dan disimpan!")

    with tab_up3:
      st.subheader("3. Upload Data Siswa Remedial dari Guru")
      st.markdown(
          "Upload data remedial dari guru mapel. Nilai siswa yang masuk daftar"
          " ini akan otomatis di-reset menjadi `0` (status remedial) agar"
          " wali kelas dan siswa tahu bahwa mereka harus mengikuti remedial."
      )
      file_rem = st.file_uploader(
          "Pilih file Data Remedial", type=["xlsx", "csv"], key="remedial_file"
      )
      if file_rem:
        df_r = (
            pd.read_csv(file_rem)
            if file_rem.name.endswith(".csv")
            else pd.read_excel(file_rem)
        )
        df_r.columns = [str(col).strip() for col in df_r.columns]
        col_map_r = {c.lower(): c for c in df_r.columns}
        tes_rc = col_map_r.get("nama tes") or col_map_r.get("namates")
        nis_rc = (
            col_map_r.get("username")
            or col_map_r.get("nis")
            or col_map_r.get("no induk")
        )
        nama_rc = col_map_r.get("nama")
        kelas_rc = col_map_r.get("group") or col_map_r.get("kelas")

        if tes_rc and nama_rc:
          st.markdown("**Pratinjau Data Remedial:**")
          st.dataframe(
              df_r.head(), hide_index=True, use_container_width=True
          )
          if st.button("Simpan Data Remedial & Reset Nilai ke 0 (Remedial)"):
            cursor = conn.cursor()
            count_saved = 0
            for _, row in df_r.iterrows():
              nama_val = str(row[nama_rc]).strip()
              tes_val = str(row[tes_rc]).strip()
              clean_r_tes = clean_tes_name(tes_val)

              nis_val = None
              if nis_rc and pd.notna(row[nis_rc]):
                nis_raw = row[nis_rc]
                nis_val = (
                    str(int(nis_raw))
                    if isinstance(nis_raw, float)
                    else str(nis_raw).strip()
                )

              kelas_val = (
                  str(row[kelas_rc]).strip()
                  if kelas_rc and pd.notna(row[kelas_rc])
                  else "-"
              )

              # Auto-Fill NIS & Kelas dari Master jika kosong/nan
              cursor.execute(
                  "SELECT nis, kelas FROM master_siswa WHERE LOWER(nama) ="
                  " LOWER(?)",
                  (nama_val,),
              )
              res_m = cursor.fetchone()
              if res_m:
                if not nis_val or nis_val.lower() == "nan":
                  nis_val = res_m[0]
                if kelas_val == "-" or not kelas_val:
                  kelas_val = res_m[1]

              if not nis_val or nis_val.lower() == "nan":
                nis_val = "UNKNOWN_" + nama_val

              # 1. Simpan ke remedial_siswa
              cursor.execute(
                  """
                                INSERT OR REPLACE INTO remedial_siswa (nama_tes, nis, nama, kelas)
                                VALUES (?, ?, ?, ?)
                            """,
                  (tes_val, nis_val, nama_val, kelas_val),
              )

              # 2. Cari test name yang tepat di hasil_ujian jika ada
              cursor.execute(
                  "SELECT id, nama_tes FROM hasil_ujian WHERE nis = ?",
                  (nis_val,),
              )
              all_h = cursor.fetchall()
              target_tes = tes_val
              for _, h_tes in all_h:
                if clean_tes_name(h_tes) == clean_r_tes:
                  target_tes = h_tes
                  break

              # 3. Replace nilai di hasil_ujian menjadi 0 dengan status 'remedial'
              cursor.execute(
                  """
                                INSERT OR REPLACE INTO hasil_ujian (nama_tes, nis, nama, kelas, poin, status)
                                VALUES (?, ?, ?, ?, 0.0, 'remedial')
                            """,
                  (target_tes, nis_val, nama_val, kelas_val),
              )
              count_saved += 1

            conn.commit()
            st.success(
                f"Berhasil menyimpan {count_saved} data remedial. Nilai siswa"
                " tersebut telah di-reset menjadi 0 (status remedial)."
            )

    with tab_up4:
      st.subheader("4. Update Nilai & Selesaikan Remedial")
      file_up_rem = st.file_uploader(
          "Pilih file Update Nilai Remedial / Susulan",
          type=["xlsx", "csv"],
          key="up_rem_score",
      )
      if file_up_rem:
        df_up_r = (
            pd.read_csv(file_up_rem)
            if file_up_rem.name.endswith(".csv")
            else pd.read_excel(file_up_rem)
        )
        df_up_r.columns = [str(col).strip() for col in df_up_r.columns]
        col_map_ur = {c.lower(): c for c in df_up_r.columns}
        tes_ur = col_map_ur.get("nama tes") or col_map_ur.get("namates")
        nis_ur = (
            col_map_ur.get("username")
            or col_map_ur.get("nis")
            or col_map_ur.get("no induk")
        )
        nama_ur = col_map_ur.get("nama")
        poin_ur = col_map_ur.get("poin") or col_map_ur.get("nilai")

        if tes_ur and (nis_ur or nama_ur) and poin_ur:
          st.markdown("**Pratinjau Update Nilai:**")
          st.dataframe(
              df_up_r.head(), hide_index=True, use_container_width=True
          )
          if st.button("🚀 Proses Update Nilai Remedial"):
            cursor = conn.cursor()
            updated_count = 0
            for _, row in df_up_r.iterrows():
              poin_val = parse_poin(row[poin_ur])
              tes_val = str(row[tes_ur]).strip()
              clean_up_tes = clean_tes_name(tes_val)

              nama_val = (
                  str(row[nama_ur]).strip()
                  if nama_ur and pd.notna(row[nama_ur])
                  else None
              )
              nis_val = None
              if nis_ur and pd.notna(row[nis_ur]):
                nis_raw = row[nis_ur]
                nis_val = (
                    str(int(nis_raw))
                    if isinstance(nis_raw, float)
                    else str(nis_raw).strip()
                )

              if not nis_val and nama_val:
                cursor.execute(
                    "SELECT nis, kelas FROM master_siswa WHERE LOWER(nama) ="
                    " LOWER(?)",
                    (nama_val,),
                )
                res_m = cursor.fetchone()
                if res_m:
                  nis_val, kelas_val = res_m[0], res_m[1]
                else:
                  kelas_val = "-"
                  nis_val = "UNKNOWN_" + nama_val
              elif nis_val and not nama_val:
                cursor.execute(
                    "SELECT nama, kelas FROM master_siswa WHERE nis = ?",
                    (nis_val,),
                )
                res_m = cursor.fetchone()
                if res_m:
                  nama_val, kelas_val = res_m[0], res_m[1]
                else:
                  nama_val = "Siswa"
                  kelas_val = "-"
              else:
                kelas_val = "-"

              # Update ke hasil_ujian dengan status remedial
              cursor.execute(
                  """
                                INSERT OR REPLACE INTO hasil_ujian (nama_tes, nis, nama, kelas, poin, status)
                                VALUES (?, ?, ?, ?, ?, 'remedial')
                            """,
                  (tes_val, nis_val, nama_val, kelas_val, poin_val),
              )
              updated_count += cursor.rowcount

            conn.commit()
            st.success(
                f"Berhasil memperbarui {updated_count} nilai siswa remedial!"
            )

    with tab_up5:
      st.subheader("5. Kelola & Hapus Data di Database")
      df_mapel_del = pd.read_sql(
          "SELECT DISTINCT nama_tes FROM hasil_ujian", conn
      )
      if df_mapel_del.empty:
        st.info("Belum ada data hasil ujian di dalam database.")
      else:
        pilih_hapus_tes_list = st.multiselect(
            "Pilih Mata Pelajaran Ujian yang ingin dihapus:",
            df_mapel_del["nama_tes"].tolist(),
            key="del_ujian_multiselect",
        )
        if st.button("🗑️ Hapus Data Ujian yang Dipilih", type="primary"):
          if pilih_hapus_tes_list:
            cursor = conn.cursor()
            for tes in pilih_hapus_tes_list:
              cursor.execute(
                  "DELETE FROM hasil_ujian WHERE nama_tes = ?", (tes,)
              )
            conn.commit()
            st.success("Data ujian berhasil dihapus!")
            st.rerun()

      st.markdown("---")
      col_r1, col_r2, col_r3 = st.columns(3)
      with col_r1:
        if st.button("⚠️ Kosongkan SEMUA Hasil Ujian"):
          cursor = conn.cursor()
          cursor.execute("DELETE FROM hasil_ujian")
          conn.commit()
          st.warning("Semua data hasil ujian dikosongkan!")
          st.rerun()
      with col_r2:
        if st.button("⚠️ Kosongkan SEMUA Data Master"):
          cursor = conn.cursor()
          cursor.execute("DELETE FROM master_siswa")
          conn.commit()
          st.warning("Semua data master dikosongkan!")
          st.rerun()
      with col_r3:
        if st.button("⚠️ Kosongkan SEMUA Data Remedial"):
          cursor = conn.cursor()
          cursor.execute("DELETE FROM remedial_siswa")
          conn.commit()
          st.warning("Semua data remedial dikosongkan!")
          st.rerun()

  # --- MENU 2: GURU PREVIEW & DOWNLOAD REKAP ---
  elif menu == "2. Guru: Download Rekap Nilai":
    st.header("Dashboard Guru: Pratinjau & Download Rekap Nilai Per Kelas")

    df_mapel = pd.read_sql("SELECT DISTINCT nama_tes FROM hasil_ujian", conn)
    df_kelas = pd.read_sql("SELECT DISTINCT kelas FROM master_siswa", conn)

    if df_mapel.empty or df_kelas.empty:
      st.warning(
          "Data belum lengkap. Operator harus meng-upload Data Master Siswa dan"
          " Hasil Ujian terlebih dahulu."
      )
    else:
      pilih_mapel = st.selectbox(
          "Pilih Mata Pelajaran / Nama Tes:", df_mapel["nama_tes"].tolist()
      )

      list_kelas_tersedia = df_kelas["kelas"].tolist()
      pilih_kelas = st.multiselect(
          "Pilih Kelas (Bisa pilih lebih dari 1 kelas sekaligus):",
          list_kelas_tersedia,
      )

      if pilih_kelas:
        st.markdown("---")
        st.subheader("👀 Pratinjau (Preview) Data Nilai")

        preview_data = {}
        cleaned_selected_tes = clean_tes_name(pilih_mapel)

        df_all_ujian = pd.read_sql(
            "SELECT nis, nama, poin, status, nama_tes FROM hasil_ujian", conn
        )
        df_matched_ujian = df_all_ujian[
            df_all_ujian["nama_tes"].apply(clean_tes_name)
            == cleaned_selected_tes
        ]

        dict_nilai = {}
        for _, n_row in df_matched_ujian.iterrows():
          dict_nilai[str(n_row["nis"])] = (n_row["poin"], n_row["status"])
          dict_nilai[clean_name(n_row["nama"])] = (
              n_row["poin"],
              n_row["status"],
          )

        for kelas in pilih_kelas:
          st.markdown(f"**Kelas: {kelas}**")
          df_m_kelas = pd.read_sql(
              "SELECT nis, nama FROM master_siswa WHERE kelas = ?",
              conn,
              params=(kelas,),
          )

          rows = []
          for idx, (_, s_row) in enumerate(df_m_kelas.iterrows(), 1):
            nis_siswa = str(s_row["nis"])
            nama_siswa = s_row["nama"]
            clean_n_siswa = clean_name(nama_siswa)

            data_n = dict_nilai.get(nis_siswa) or dict_nilai.get(
                clean_n_siswa, (0.0, "normal")
            )
            nilai_pg = data_n[0]
            status_n = data_n[1]

            if status_n == "remedial":
              ket = "Sudah Remedial"
            elif status_n == "susulan":
              ket = "Sudah Susulan"
            else:
              ket = "-"

            rows.append({
                "NO": idx,
                "NAMA": nama_siswa,
                "NILAI PG": nilai_pg,
                "KET": ket,
            })

          df_preview = pd.DataFrame(rows)
          if df_preview.empty:
            st.info(f"Tidak ada data siswa untuk kelas {kelas}.")
          else:
            st.dataframe(
                df_preview, hide_index=True, use_container_width=True
            )

          preview_data[kelas] = df_preview

        st.markdown("---")
        if st.button("📥 Generate & Download Excel (Multi-Sheet)"):
          output = io.BytesIO()
          wb = openpyxl.Workbook()
          default_sheet = wb.active
          wb.remove(default_sheet)

          font_title = Font(name="Calibri", size=11, bold=True)
          font_bold = Font(name="Calibri", size=11, bold=True)
          fill_header = PatternFill(
              start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
          )
          thin_border = Border(
              left=Side(style="thin", color="000000"),
              right=Side(style="thin", color="000000"),
              top=Side(style="thin", color="000000"),
              bottom=Side(style="thin", color="000000"),
          )

          for kelas, df_prev in preview_data.items():
            safe_sheet_name = re.sub(r"[\\/?:*\[\]]", "-", str(kelas))[:31]
            ws = wb.create_sheet(title=safe_sheet_name)

            ws.merge_cells("A1:D1")
            ws["A1"] = f"MATA PELAJARAN : {pilih_mapel}"
            ws["A1"].font = font_title
            ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

            ws.merge_cells("A2:D2")
            ws["A2"] = f"KELAS : {kelas}"
            ws["A2"].font = font_title
            ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

            headers = ["NO", "NAMA", "NILAI PG", "KET"]
            for col_num, h_title in enumerate(headers, 1):
              cell = ws.cell(row=4, column=col_num, value=h_title)
              cell.font = font_bold
              cell.fill = fill_header
              cell.alignment = Alignment(horizontal="center", vertical="center")
              cell.border = thin_border

            for _, r_data in df_prev.iterrows():
              r_idx = int(r_data["NO"])
              row_num = 4 + r_idx

              c_no = ws.cell(row=row_num, column=1, value=r_idx)
              c_nama = ws.cell(row=row_num, column=2, value=r_data["NAMA"])
              c_nilai = ws.cell(row=row_num, column=3, value=r_data["NILAI PG"])
              c_ket = ws.cell(row=row_num, column=4, value=r_data["KET"])

              nis_s = str(
                  pd.read_sql(
                      "SELECT nis FROM master_siswa WHERE kelas = ? AND nama ="
                      " ?",
                      conn,
                      params=(kelas, r_data["NAMA"]),
                  ).iloc[0]["nis"]
              )
              data_s = dict_nilai.get(nis_s) or dict_nilai.get(
                  clean_name(r_data["NAMA"]), (0.0, "normal")
              )
              status_s = data_s[1]

              if status_s == "remedial":
                c_nilai.font = Font(
                    name="Calibri", size=11, color="008000", bold=True
                )
              elif status_s == "susulan":
                c_nilai.font = Font(
                    name="Calibri", size=11, color="0000FF", bold=True
                )
              else:
                c_nilai.font = Font(name="Calibri", size=11)

              c_no.alignment = Alignment(
                  horizontal="center", vertical="center"
              )
              c_nama.alignment = Alignment(horizontal="left", vertical="center")
              c_nilai.alignment = Alignment(
                  horizontal="center", vertical="center"
              )
              c_ket.alignment = Alignment(
                  horizontal="center", vertical="center"
              )

              c_no.border = thin_border
              c_nama.border = thin_border
              c_nilai.border = thin_border
              c_ket.border = thin_border

            ws.column_dimensions["A"].width = 8
            ws.column_dimensions["B"].width = 38
            ws.column_dimensions["C"].width = 16
            ws.column_dimensions["D"].width = 20

          wb.save(output)
          output.seek(0)

          st.success("File rekap multi-kelas berhasil dibuat!")
          st.download_button(
              label="📥 Download File Excel Rekap (Multi-Sheet)",
              data=output,
              file_name=f"Rekap_Nilai_{pilih_mapel}.xlsx",
              mime=(
                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
              ),
          )

  # --- MENU 3: CEK SISWA BELUM UJIAN ---
  elif menu == "3. Cek Siswa Belum Ujian":
    st.header("Pelacak Siswa Belum Mengikuti Ujian")

    df_mapel = pd.read_sql("SELECT DISTINCT nama_tes FROM hasil_ujian", conn)
    if df_mapel.empty:
      st.warning("Belum ada data ujian yang di-upload.")
    else:
      pilih_mapel = st.selectbox(
          "Pilih Mata Pelajaran:",
          df_mapel["nama_tes"].tolist(),
          key="cek_mapel",
      )

      cleaned_cek_tes = clean_tes_name(pilih_mapel)
      df_all_ujian = pd.read_sql(
          "SELECT nis, nama, nama_tes FROM hasil_ujian", conn
      )
      matched_nis = set()
      for _, h in df_all_ujian.iterrows():
        if clean_tes_name(h["nama_tes"]) == cleaned_cek_tes:
          matched_nis.add(str(h["nis"]))
          matched_nis.add(clean_name(h["nama"]))

      df_master_all = pd.read_sql(
          "SELECT kelas, nis, nama FROM master_siswa ORDER BY kelas, nama", conn
      )
      belum_rows = []
      for _, m in df_master_all.iterrows():
        if (
            str(m["nis"]) not in matched_nis
            and clean_name(m["nama"]) not in matched_nis
        ):
          belum_rows.append(m)

      df_belum = pd.DataFrame(belum_rows)

      if df_belum.empty:
        st.info("Hebat! Semua siswa sudah mengikuti ujian untuk mata pelajaran ini.")
      else:
        list_kelas_belum = df_belum["kelas"].unique().tolist()
        select_all_kls_blm = st.checkbox(
            "Pilih Semua Kelas", value=True, key="select_all_kls_blm_chk"
        )
        if select_all_kls_blm:
          pilih_kelas_blm = st.multiselect(
              "Pilih Kelas:",
              list_kelas_belum,
              default=list_kelas_belum,
              key="multiselect_kelas_blm",
          )
        else:
          pilih_kelas_blm = st.multiselect(
              "Pilih Kelas:", list_kelas_belum, key="multiselect_kelas_blm"
          )

        if pilih_kelas_blm:
          df_filtered_belum = df_belum[df_belum["kelas"].isin(pilih_kelas_blm)]
        else:
          df_filtered_belum = pd.DataFrame(columns=df_belum.columns)

        if df_filtered_belum.empty:
          st.info("Tidak ada siswa belum ujian untuk kelas yang dipilih.")
        else:
          st.warning(
              f"Ditemukan {len(df_filtered_belum)} siswa yang belum mengikuti"
              " ujian."
          )
          st.dataframe(
              df_filtered_belum, hide_index=True, use_container_width=True
          )

          output_blm = io.BytesIO()
          with pd.ExcelWriter(output_blm, engine="openpyxl") as writer:
            df_filtered_belum.to_excel(
                writer, sheet_name="Siswa Belum Ujian", index=False
            )
          output_blm.seek(0)

          st.download_button(
              label="📥 Download Daftar Siswa Belum Ujian ke Excel",
              data=output_blm,
              file_name=(
                  f"Siswa_Belum_Ujian_{pilih_mapel.replace('/', '-')}.xlsx"
              ),
              mime=(
                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
              ),
          )

  # --- MENU 4: CEK & DOWNLOAD PESERTA REMEDIAL (AUTO-SYNC FILTER) ---
  elif menu == "4. Cek & Download Peserta Remedial":
    st.header("Pelacak & Rekap Peserta Didik Remedial")

    df_mapel_rem = pd.read_sql(
        "SELECT DISTINCT nama_tes FROM remedial_siswa", conn
    )
    df_kelas_rem = pd.read_sql(
        "SELECT DISTINCT kelas FROM remedial_siswa", conn
    )

    if df_mapel_rem.empty or df_kelas_rem.empty:
      st.warning("Belum ada data remedial yang di-upload oleh operator.")
    else:
      list_mapel_rem = df_mapel_rem["nama_tes"].tolist()
      list_kelas_rem = df_kelas_rem["kelas"].tolist()

      pilih_mapel_rem = st.multiselect(
          "Pilih Mata Pelajaran Remedial:",
          list_mapel_rem,
          key="multiselect_mapel_rem",
      )

      if "multiselect_kelas_rem" not in st.session_state:
        st.session_state.multiselect_kelas_rem = []


      def toggle_all_kelas():
        if st.session_state.get("select_all_kls_rem_chk", False):
          st.session_state.multiselect_kelas_rem = list_kelas_rem
        else:
          st.session_state.multiselect_kelas_rem = []


      st.checkbox(
          "Pilih Semua Kelas Remedial",
          key="select_all_kls_rem_chk",
          on_change=toggle_all_kelas,
      )

      pilih_kelas_rem = st.multiselect(
          "Pilih Kelas:", list_kelas_rem, key="multiselect_kelas_rem"
      )

      if pilih_mapel_rem or pilih_kelas_rem:
        query_r = "SELECT nama_tes, kelas, nis, nama FROM remedial_siswa WHERE 1=1"
        params = []

        if pilih_mapel_rem:
          placeholders_m = ",".join(["?"] * len(pilih_mapel_rem))
          query_r += f" AND nama_tes IN ({placeholders_m})"
          params.extend(pilih_mapel_rem)

        if pilih_kelas_rem:
          placeholders_k = ",".join(["?"] * len(pilih_kelas_rem))
          query_r += f" AND kelas IN ({placeholders_k})"
          params.extend(pilih_kelas_rem)

        query_r += " ORDER BY nama_tes, kelas, nama"
        df_hasil_rem = pd.read_sql(query_r, conn, params=params)

        # --- SMART SYNC: Ambil daftar siswa yang sudah ujian di hasil_ujian ---
        df_all_ujian = pd.read_sql(
            "SELECT nis, nama, nama_tes FROM hasil_ujian", conn
        )
        completed_set = set()
        for _, h in df_all_ujian.iterrows():
          c_nis = str(h["nis"]).strip()
          c_name = clean_name(h["nama"])
          c_tes = clean_tes_name(h["nama_tes"])
          completed_set.add((c_nis, c_tes))
          completed_set.add((c_name, c_tes))

        filtered_rows = []
        for _, r in df_hasil_rem.iterrows():
          r_nis = str(r["nis"]).strip()
          r_name = clean_name(r["nama"])
          r_tes = clean_tes_name(r["nama_tes"])

          is_done = False
          for comp_id, comp_tes in completed_set:
            if (comp_id == r_nis or comp_id == r_name) and (
                comp_tes == r_tes
                or comp_tes in r_tes
                or r_tes in comp_tes
                or "b. arab" in comp_tes
                and "arab" in r_tes
            ):
              is_done = True
              break
          if not is_done:
            filtered_rows.append(r)

        df_final_rem = pd.DataFrame(filtered_rows)

        if df_final_rem.empty:
          st.info(
              "Hebat! Semua siswa pada filter ini sudah menyelesaikan ujian"
              " remedial/susulan."
          )
        else:
          st.warning(
              f"Ditemukan {len(df_final_rem)} siswa yang masih berstatus"
              " remedial/susulan."
          )
          st.dataframe(
              df_final_rem, hide_index=True, use_container_width=True
          )

          output_rem = io.BytesIO()
          with pd.ExcelWriter(output_rem, engine="openpyxl") as writer:
            df_final_rem.to_excel(
                writer, sheet_name="Peserta Remedial", index=False
            )
          output_rem.seek(0)

          st.download_button(
              label="📥 Download Daftar Remedial ke Excel",
              data=output_rem,
              file_name="Daftar_Remedial_Pilihan.xlsx",
              mime=(
                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
              ),
          )
